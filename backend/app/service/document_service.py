import hashlib
import shutil
import subprocess
import tempfile
import unicodedata
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_COLOR_INDEX
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import Highlight, Text
from pypdf.generic import ArrayObject, FloatObject

from app.core.config import settings
from app.models.review import ReviewRisk, ReviewTask
from app.service.file_service import FileParseError, FileService


class DocumentProcessingError(RuntimeError):
    pass


class DocumentService:
    def __init__(self, file_service: FileService | None = None) -> None:
        self.file_service = file_service or FileService()
        self.result_root = self.file_service.storage_root / "results"
        self.result_root.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def snapshot_file(file_id: str, item: dict) -> dict:
        extension = str(item.get("file_extension") or "").lower()
        pages = item.get("pages") or []
        return {
            "file_id": file_id,
            "file_name": str(item.get("file_name") or file_id),
            "file_extension": extension,
            "mime_type": str(item.get("mime_type") or "application/octet-stream"),
            "page_count": int(item.get("page_count") or len(pages) or 1),
            "ocr_used": bool(item.get("ocr_used", False)),
            "ocr_status": str(item.get("ocr_status", "not_required")),
            "ocr_message": item.get("ocr_message"),
            "text_source": str(item.get("text_source", "native")),
            "preview_available": bool(item.get("preview_available", False)),
            "pagination_source": str(item.get("pagination_source", "logical")),
            "office_text_source": item.get("office_text_source"),
            "preview_quality": item.get("preview_quality"),
            "pages": pages,
        }

    def documents(self, task: ReviewTask) -> list[dict]:
        if task.document_files:
            documents = []
            for item in task.document_files:
                normalized = dict(item)
                normalized["pages"] = self._normalize_pages(item.get("pages", []))
                normalized["page_count"] = len(normalized["pages"]) or 1
                documents.append(normalized)
            return documents
        main_pages = task.document_pages or self.file_service._paginate_blocks(
            self.file_service._text_blocks(task.contract_text)
        )
        main = self.snapshot_file(
            "main",
            {
                "file_name": task.contract_name,
                "file_extension": Path(task.contract_name).suffix.lstrip("."),
                **(task.document_meta or {}),
                "pages": main_pages,
            },
        )
        documents = [main]
        for index, attachment in enumerate(task.attachment_texts, start=1):
            pages = attachment.get("pages") or self.file_service._paginate_blocks(
                self.file_service._text_blocks(str(attachment.get("content", "")))
            )
            documents.append(
                self.snapshot_file(
                    f"attachment:{index}", {**attachment, "pages": pages}
                )
            )
        return documents

    def document(self, task: ReviewTask, file_id: str = "main") -> dict:
        document = next(
            (item for item in self.documents(task) if item.get("file_id") == file_id),
            None,
        )
        if not document:
            raise FileParseError("合同材料不存在。")
        return document

    def pages(self, task: ReviewTask, file_id: str = "main") -> list[dict]:
        document = self.document(task, file_id) if task.document_files else None
        if document:
            return document["pages"]
        if file_id == "main" and task.document_pages:
            return self._normalize_pages(task.document_pages)
        recovered, _meta = self._recover_document(task, file_id=file_id)
        return recovered

    def ensure_document_snapshot(self, task: ReviewTask) -> bool:
        current_pages = task.document_pages or []
        needs_recovery = not current_pages
        meta = dict(task.document_meta or {})
        extension = Path(task.contract_name).suffix.lower().lstrip(".") or "txt"
        needs_office_recovery = (
            extension in {"doc", "docx"}
            and (
                meta.get("preview_quality") == "unusable"
                or meta.get("office_text_source") != "structured"
                or not self._snapshot_matches_contract_text(
                    current_pages, task.contract_text
                )
            )
        )
        has_page_text = any(str(page.get("text", "")).strip() for page in current_pages)
        needs_pdf_recovery = extension == "pdf" and (
            any(
                str(page.get("text", "")).strip()
                and not self.file_service._text_is_usable(str(page.get("text", "")))
                for page in current_pages
            )
            or (
                not has_page_text
                and meta.get("ocr_status") not in {"failed", "unavailable"}
            )
        )
        if current_pages:
            needs_recovery = any(
                not block.get("block_id")
                for page in current_pages
                for block in page.get("blocks", [])
            )
        if needs_recovery or needs_office_recovery or needs_pdf_recovery:
            pages, recovered_meta = self._recover_document(
                task, prefer_source=needs_office_recovery or needs_pdf_recovery
            )
        else:
            pages = self._normalize_pages(current_pages)
            recovered_meta = {}

        changed = pages != current_pages
        if changed:
            task.document_pages = pages
        preview_quality = recovered_meta.get("preview_quality", meta.get("preview_quality"))
        if extension in {"doc", "docx"} and not preview_quality:
            preview_quality = self._existing_office_preview_quality(task)
        desired_meta = {
            **recovered_meta,
            "file_name": str(meta.get("file_name") or task.contract_name),
            "file_extension": str(meta.get("file_extension") or extension),
            "mime_type": str(
                meta.get("mime_type") or self.file_service._mime_type(extension)
            ),
            "page_count": len(pages) or 1,
            "ocr_used": bool(recovered_meta.get("ocr_used", meta.get("ocr_used", False))),
            "ocr_status": str(
                recovered_meta.get("ocr_status", meta.get("ocr_status", "not_required"))
            ),
            "ocr_message": recovered_meta.get("ocr_message", meta.get("ocr_message")),
            "text_source": str(
                recovered_meta.get("text_source", meta.get("text_source", "native"))
            ),
            "preview_available": bool(
                preview_quality == "usable"
                if extension in {"doc", "docx"}
                else recovered_meta.get(
                    "preview_available",
                    meta.get("preview_available", extension == "pdf"),
                )
            ),
            "pagination_source": str(
                recovered_meta.get(
                    "pagination_source",
                    meta.get(
                        "pagination_source",
                        "physical_pdf" if extension == "pdf" else "logical",
                    ),
                )
            ),
            "office_text_source": (
                str(recovered_meta.get("office_text_source", "structured"))
                if extension in {"doc", "docx"}
                else None
            ),
            "preview_quality": preview_quality,
        }
        if desired_meta != meta:
            task.document_meta = desired_meta
            changed = True
        existing_files = {
            str(item.get("file_id")): item for item in (task.document_files or [])
        }
        material_files = [
            self.snapshot_file("main", {**desired_meta, "pages": pages})
        ]
        for index, attachment in enumerate(task.attachment_texts, start=1):
            file_id = f"attachment:{index}"
            existing = existing_files.get(file_id)
            if existing and existing.get("pages"):
                item = dict(existing)
                item["pages"] = self._normalize_pages(existing["pages"])
            else:
                recovered_pages, recovered_meta = self._recover_document(
                    task, file_id=file_id, prefer_source=True
                )
                item = self.snapshot_file(
                    file_id,
                    {
                        **attachment,
                        **recovered_meta,
                        "file_name": attachment.get("file_name")
                        or task.attachment_names[index - 1],
                        "pages": recovered_pages,
                    },
                )
            item["page_count"] = len(item.get("pages", [])) or 1
            material_files.append(item)
        if material_files != (task.document_files or []):
            task.document_files = material_files
            changed = True
        return changed

    def _existing_office_preview_quality(self, task: ReviewTask) -> str:
        if not task.user_id:
            return "unavailable"
        try:
            manifest = self.file_service.load_manifest(task.upload_id, task.user_id)
        except (FileParseError, OSError, ValueError):
            return "unavailable"
        main_file = manifest["main_file"]
        source = Path(str(main_file.get("stored_path") or ""))
        candidates = [
            Path(str(main_file.get("preview_path") or "")),
            source.with_name(f"{source.stem}-preview.pdf"),
            self.result_root / str(task.id) / "original-preview.pdf",
        ]
        preview = next((path for path in candidates if path.is_file()), None)
        if not preview:
            return "unavailable"
        return (
            "usable"
            if self.file_service._office_preview_is_usable(preview)
            else "unusable"
        )

    def ensure_risk_locations(self, task: ReviewTask, risks: list[ReviewRisk]) -> bool:
        changed = False
        for risk in risks:
            if self._stored_location_is_valid(task, risk):
                continue
            hint = risk.location_hint or {}
            location = self.locate_text(
                task,
                risk.source,
                risk.original_text,
                source_file_id=risk.source_file_id,
                page_number=hint.get("page_number"),
                context_before=hint.get("context_before"),
                context_after=hint.get("context_after"),
            )
            file_id = risk.source_file_id or self._source_file_id(task, risk.source)
            primary = (location or {}).get("location_spans", [{}])[0]
            values = {
                "source_file_id": file_id,
                "page_number": primary.get("page_number") if location else None,
                "block_id": primary.get("block_id") if location else None,
                "start_offset": primary.get("start_offset") if location else None,
                "end_offset": primary.get("end_offset") if location else None,
                "anchor_text": (
                    primary.get("anchor_text") if location else risk.original_text
                ),
                "anchor_hash": primary.get("anchor_hash") if location else None,
                "location_spans": (
                    location.get("location_spans", []) if location else []
                ),
                "locatable": bool(location),
            }
            for field, value in values.items():
                if getattr(risk, field) != value:
                    setattr(risk, field, value)
                    changed = True
        return changed

    def locate_text(
        self,
        task: ReviewTask,
        source: str | None,
        original_text: str | None,
        *,
        source_file_id: str | None = None,
        page_number: int | None = None,
        context_before: str | None = None,
        context_after: str | None = None,
    ) -> dict | None:
        if not original_text:
            return None
        if len(self._normalize(original_text)) < 4:
            return None
        file_id = source_file_id or self._source_file_id(task, source)
        if not file_id:
            return None
        try:
            document = self.document(task, file_id)
        except FileParseError:
            return None
        if source and source not in {document.get("file_name"), "需人工核验"}:
            return None

        haystack: list[str] = []
        references: list[tuple[int, str, int, str]] = []
        for page in document.get("pages", []):
            current_page = int(page.get("page_number", 1))
            for block in page.get("blocks", []):
                block_id = str(block.get("block_id", ""))
                block_text = str(block.get("text", ""))
                for character, offset in self._normalized_with_offsets(block_text):
                    haystack.append(character)
                    references.append((current_page, block_id, offset, block_text))

        quote = self._normalized_value(original_text)
        if not quote:
            return None
        text = "".join(haystack)
        candidates: list[dict] = []
        cursor = 0
        while True:
            found = text.find(quote, cursor)
            if found < 0:
                break
            end = found + len(quote)
            spans = self._spans_from_references(file_id, references[found:end])
            if spans:
                candidates.append(
                    {
                        **spans[0],
                        "location_spans": spans,
                        "normalized_start": found,
                        "normalized_end": end,
                    }
                )
            cursor = found + max(1, len(quote))

        # A model-provided page number is advisory. Office pagination can differ from
        # the model's plain-text page estimate, so a unique exact quote must win.
        if len(candidates) > 1 and page_number:
            page_candidates = [
                item for item in candidates if item["page_number"] == page_number
            ]
            if page_candidates:
                candidates = page_candidates
        before = self._normalized_value(context_before or "")
        after = self._normalized_value(context_after or "")
        if before:
            context_candidates = [
                item
                for item in candidates
                if text[: item["normalized_start"]].endswith(before)
            ]
            if context_candidates:
                candidates = context_candidates
        if after:
            context_candidates = [
                item
                for item in candidates
                if text[item["normalized_end"] :].startswith(after)
            ]
            if context_candidates:
                candidates = context_candidates
        if len(candidates) != 1:
            return None
        result = dict(candidates[0])
        result.pop("normalized_start", None)
        result.pop("normalized_end", None)
        return result

    def get_page(
        self, task: ReviewTask, page_number: int, file_id: str = "main"
    ) -> dict:
        pages = self.pages(task, file_id)
        if page_number < 1 or page_number > len(pages):
            raise FileParseError("文档页码不存在。")
        return pages[page_number - 1]

    def render_page_png(
        self, task: ReviewTask, user_id: int, page_number: int, file_id: str = "main"
    ) -> bytes:
        pdf_path = self._ensure_preview_pdf(task, user_id, file_id)
        pdftoppm = shutil.which("pdftoppm")
        if not pdftoppm:
            raise DocumentProcessingError("当前环境未安装 PDF 页面渲染组件。")
        with tempfile.TemporaryDirectory(prefix="contract-preview-") as temp_dir:
            prefix = Path(temp_dir) / "page"
            result = subprocess.run(
                [
                    pdftoppm,
                    "-f",
                    str(page_number),
                    "-l",
                    str(page_number),
                    "-singlefile",
                    "-png",
                    "-r",
                    "130",
                    str(pdf_path),
                    str(prefix),
                ],
                capture_output=True,
                check=False,
                timeout=settings.ocr_timeout_seconds,
            )
            image_path = prefix.with_suffix(".png")
            if result.returncode != 0 or not image_path.is_file():
                raise DocumentProcessingError("文档页面预览生成失败。")
            return image_path.read_bytes()

    def generate_marked_pdf(
        self, task: ReviewTask, user_id: int, risks: list[ReviewRisk]
    ) -> Path:
        output_dir = self.result_root / str(task.id)
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / "review-marked.pdf"
        source_pdf = self._ensure_preview_pdf(task, user_id)
        reader = PdfReader(source_pdf)
        writer = PdfWriter()
        writer.clone_document_from_reader(reader)
        fragments_by_page = [self._text_fragments(page) for page in reader.pages]
        color_map = {"high": "d1493f", "medium": "d99020", "low": "4f8a61"}

        for risk in risks:
            located_page = self.risk_page_number(task, risk)
            page_number = located_page or (1 if writer.pages else None)
            if not page_number or page_number > len(writer.pages):
                continue
            page_index = page_number - 1
            matched = False
            quote = (
                self._normalize(risk.original_text or "")
                if risk.source_file_id in {None, "main"}
                else ""
            )
            for fragment_text, rect in fragments_by_page[page_index]:
                fragment = self._normalize(fragment_text)
                if not quote or not fragment or min(len(fragment), len(quote)) < 3:
                    continue
                if fragment in quote or quote in fragment:
                    writer.add_annotation(
                        page_number=page_index,
                        annotation=Highlight(
                            rect=rect,
                            quad_points=self._quad_points(rect),
                            highlight_color=color_map.get(risk.risk_level, "d1493f"),
                            printing=True,
                        ),
                    )
                    matched = True
            page = writer.pages[page_index]
            width = float(page.mediabox.width)
            height = float(page.mediabox.height)
            note_y = max(24.0, height - 36.0 - (risk.id % 12) * 18.0)
            writer.add_annotation(
                page_number=page_index,
                annotation=Text(
                    rect=(max(12.0, width - 32.0), note_y, max(28.0, width - 12.0), note_y + 16.0),
                    text=(
                        f"{risk.risk_name}（{'已定位原文' if matched else '需人工核验'}）\n"
                        f"来源：{risk.source or '未提供'}\n"
                        f"风险说明：{risk.description}\n修改建议：{risk.suggestion}"
                    ),
                    open=False,
                ),
            )
        with tempfile.NamedTemporaryFile(
            dir=output_dir, prefix="review-marked-", suffix=".pdf", delete=False
        ) as output:
            temporary_path = Path(output.name)
            writer.write(output)
        temporary_path.replace(output_path)
        return output_path

    def generate_marked_docx(
        self, task: ReviewTask, user_id: int, risks: list[ReviewRisk]
    ) -> Path:
        manifest = self.file_service.load_manifest(task.upload_id, user_id)
        source = Path(str(manifest["main_file"].get("stored_path") or ""))
        if not source.is_file() or source.suffix.lower() != ".docx":
            raise DocumentProcessingError("合同原始 DOCX 文件不存在，无法生成原格式标记版。")

        output_dir = self.result_root / str(task.id)
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / "review-marked.docx"
        document = Document(source)
        block_targets = self._map_docx_blocks(task, document)
        manual_risks: list[ReviewRisk] = []
        color_map = {
            "high": WD_COLOR_INDEX.PINK,
            "medium": WD_COLOR_INDEX.YELLOW,
            "low": WD_COLOR_INDEX.BRIGHT_GREEN,
        }

        for risk in risks:
            spans = list(risk.location_spans or [])
            if not spans and self._stored_location_is_valid(task, risk):
                spans = [
                    {
                        "source_file_id": risk.source_file_id,
                        "page_number": risk.page_number,
                        "block_id": risk.block_id,
                        "start_offset": risk.start_offset,
                        "end_offset": risk.end_offset,
                    }
                ]
            marked_runs: list[Run] = []
            reliable = bool(risk.locatable and spans)
            for span in spans:
                if span.get("source_file_id") != "main":
                    reliable = False
                    break
                target = block_targets.get(str(span.get("block_id") or ""))
                start = span.get("start_offset")
                end = span.get("end_offset")
                if not target or not isinstance(start, int) or not isinstance(end, int):
                    reliable = False
                    break
                runs = self._highlight_docx_target(
                    target,
                    start,
                    end,
                    color_map.get(risk.risk_level, WD_COLOR_INDEX.YELLOW),
                )
                if not runs:
                    reliable = False
                    break
                marked_runs.extend(runs)
            if not reliable or not marked_runs:
                manual_risks.append(risk)
                continue
            try:
                document.add_comment(
                    marked_runs,
                    text=(
                        f"{risk.risk_name}\n风险说明：{risk.description}\n"
                        f"修改建议：{risk.suggestion}"
                    ),
                    author="合同审查助手",
                    initials="AI",
                )
            except (ValueError, TypeError):
                manual_risks.append(risk)

        if manual_risks:
            document.add_page_break()
            document.add_heading("人工核验事项", level=1)
            table = document.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            for cell, title in zip(
                table.rows[0].cells,
                ("风险项", "来源", "风险说明", "修改建议"),
                strict=True,
            ):
                cell.text = title
            for risk in manual_risks:
                cells = table.add_row().cells
                cells[0].text = risk.risk_name
                cells[1].text = risk.source or "需人工核验"
                cells[2].text = risk.description
                cells[3].text = risk.suggestion

        with tempfile.NamedTemporaryFile(
            dir=output_dir, prefix="review-marked-", suffix=".docx", delete=False
        ) as output:
            temporary_path = Path(output.name)
        try:
            document.save(temporary_path)
            temporary_path.replace(output_path)
        finally:
            temporary_path.unlink(missing_ok=True)
        return output_path

    def _map_docx_blocks(
        self, task: ReviewTask, document: DocumentObject
    ) -> dict[str, dict]:
        targets: list[dict] = []
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                text = item.text.strip()
                if text:
                    targets.append(
                        {"type": "paragraph", "text": text, "paragraph": item}
                    )
            elif isinstance(item, Table):
                rows = [[cell.text.strip() for cell in row.cells] for row in item.rows]
                text = "\n".join(" | ".join(row) for row in rows)
                if text.strip():
                    targets.append(
                        {"type": "table", "text": text, "table": item, "rows": rows}
                    )

        snapshot_blocks = [
            block
            for page in self.pages(task, "main")
            for block in page.get("blocks", [])
        ]
        mapped: dict[str, dict] = {}
        used: set[int] = set()
        cursor = 0
        for block in snapshot_blocks:
            normalized = self._normalize(str(block.get("text", "")))
            candidates = [
                (index, target)
                for index, target in enumerate(targets)
                if index >= cursor
                and index not in used
                and target["type"] == block.get("type")
                and self._normalize(target["text"]) == normalized
            ]
            if not candidates:
                continue
            index, target = candidates[0]
            used.add(index)
            cursor = index + 1
            mapped[str(block.get("block_id") or "")] = target
        return mapped

    def _highlight_docx_target(
        self,
        target: dict,
        start: int,
        end: int,
        color: WD_COLOR_INDEX,
    ) -> list[Run]:
        if not (0 <= start < end <= len(str(target.get("text", "")))):
            return []
        if target["type"] == "paragraph":
            paragraph = target["paragraph"]
            leading = len(paragraph.text) - len(paragraph.text.lstrip())
            return self._highlight_paragraph_range(
                paragraph, start + leading, end + leading, color
            )

        marked: list[Run] = []
        offset = 0
        table: Table = target["table"]
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                cell_start = offset
                cell_end = cell_start + len(cell_text)
                overlap_start = max(start, cell_start)
                overlap_end = min(end, cell_end)
                if overlap_start < overlap_end:
                    marked.extend(
                        self._highlight_cell_range(
                            cell,
                            overlap_start - cell_start,
                            overlap_end - cell_start,
                            color,
                        )
                    )
                offset = cell_end
                if cell_index < len(row.cells) - 1:
                    offset += 3
            if row_index < len(table.rows) - 1:
                offset += 1
        return marked

    def _highlight_cell_range(
        self, cell: _Cell, start: int, end: int, color: WD_COLOR_INDEX
    ) -> list[Run]:
        marked: list[Run] = []
        offset = 0
        for index, paragraph in enumerate(cell.paragraphs):
            paragraph_start = offset
            paragraph_end = paragraph_start + len(paragraph.text)
            overlap_start = max(start, paragraph_start)
            overlap_end = min(end, paragraph_end)
            if overlap_start < overlap_end:
                marked.extend(
                    self._highlight_paragraph_range(
                        paragraph,
                        overlap_start - paragraph_start,
                        overlap_end - paragraph_start,
                        color,
                    )
                )
            offset = paragraph_end + (1 if index < len(cell.paragraphs) - 1 else 0)
        return marked

    @staticmethod
    def _highlight_paragraph_range(
        paragraph: Paragraph,
        start: int,
        end: int,
        color: WD_COLOR_INDEX,
    ) -> list[Run]:
        marked: list[Run] = []
        offset = 0
        for run in list(paragraph.runs):
            text = run.text
            run_start, run_end = offset, offset + len(text)
            offset = run_end
            overlap_start = max(start, run_start)
            overlap_end = min(end, run_end)
            if overlap_start >= overlap_end or not text:
                continue
            local_start = overlap_start - run_start
            local_end = overlap_end - run_start
            if local_start == 0 and local_end == len(text):
                run.font.highlight_color = color
                marked.append(run)
                continue
            pieces = [
                (text[:local_start], False),
                (text[local_start:local_end], True),
                (text[local_end:], False),
            ]
            parent = run._r.getparent()
            insert_at = parent.index(run._r)
            for piece, highlighted in pieces:
                if not piece:
                    continue
                element = deepcopy(run._r)
                cloned = Run(element, paragraph)
                cloned.text = piece
                if highlighted:
                    cloned.font.highlight_color = color
                    marked.append(cloned)
                parent.insert(insert_at, element)
                insert_at += 1
            parent.remove(run._r)
        return marked

    def _ensure_preview_pdf(
        self, task: ReviewTask, user_id: int, file_id: str = "main"
    ) -> Path:
        manifest = self.file_service.load_manifest(task.upload_id, user_id)
        material = self._manifest_material(manifest, file_id)
        source = Path(material["stored_path"])
        if not source.is_file():
            raise DocumentProcessingError("合同原始文件不存在，无法生成预览。")
        if source.suffix.lower() == ".pdf":
            return source
        uploaded_preview = Path(str(material.get("preview_path") or ""))
        if uploaded_preview.is_file() and self.file_service._office_preview_is_usable(
            uploaded_preview
        ):
            return uploaded_preview
        adjacent_preview = source.with_name(f"{source.stem}-preview.pdf")
        if (
            adjacent_preview.is_file()
            and adjacent_preview.stat().st_mtime >= source.stat().st_mtime
            and self.file_service._office_preview_is_usable(adjacent_preview)
        ):
            return adjacent_preview
        output_dir = self.result_root / str(task.id)
        output_dir.mkdir(parents=True, exist_ok=True)
        preview_name = "original-preview.pdf" if file_id == "main" else f"{file_id.replace(':', '-')}-preview.pdf"
        preview = output_dir / preview_name
        if (
            preview.is_file()
            and preview.stat().st_mtime >= source.stat().st_mtime
            and self.file_service._office_preview_is_usable(preview)
        ):
            return preview
        if not self.file_service._soffice_path():
            raise DocumentProcessingError("当前环境未安装 LibreOffice，无法保留原版式生成 PDF。")
        converted = self.file_service._convert_office_to_pdf(source)
        if not converted:
            raise DocumentProcessingError(
                "合同 PDF 转换后字体或文本校验失败，请安装中文字体后重试。"
            )
        shutil.copyfile(converted, preview)
        return preview

    @staticmethod
    def _text_fragments(page) -> list[tuple[str, tuple[float, float, float, float]]]:
        fragments: list[tuple[str, tuple[float, float, float, float]]] = []

        def visitor(text, _cm, tm, _font, font_size):
            cleaned = str(text).strip()
            if not cleaned:
                return
            x, y = float(tm[4]), float(tm[5])
            size = max(float(font_size or 10), 6.0)
            width = max(size, len(cleaned) * size * 0.72)
            fragments.append((cleaned, (x, y - size * 0.2, x + width, y + size)))

        try:
            page.extract_text(visitor_text=visitor)
        except Exception:
            return []
        return fragments

    def _find_page(self, task: ReviewTask, risk: ReviewRisk) -> int | None:
        hint = risk.location_hint or {}
        location = self.locate_text(
            task,
            risk.source,
            risk.original_text,
            source_file_id=risk.source_file_id,
            page_number=hint.get("page_number"),
            context_before=hint.get("context_before"),
            context_after=hint.get("context_after"),
        )
        return int(location["page_number"]) if location else None

    def risk_page_number(self, task: ReviewTask, risk: ReviewRisk) -> int | None:
        if risk.source_file_id not in {None, "main"}:
            return None
        if risk.locatable and risk.page_number:
            return risk.page_number
        return self._find_page(task, risk)

    def _recover_document(
        self,
        task: ReviewTask,
        prefer_source: bool = False,
        file_id: str = "main",
    ) -> tuple[list[dict], dict]:
        if task.user_id:
            try:
                manifest = self.file_service.load_manifest(task.upload_id, task.user_id)
                material = self._manifest_material(manifest, file_id)
                source = Path(str(material.get("stored_path") or ""))
                if prefer_source and source.is_file():
                    parsed = self.file_service.parse_document(
                        source,
                        material.get("file_extension") or source.suffix.lstrip("."),
                    )
                    return self._normalize_pages(parsed["pages"]), parsed
                if material.get("pages"):
                    return self._normalize_pages(material["pages"]), material
                if source.is_file():
                    parsed = self.file_service.parse_document(
                        source, material.get("file_extension") or source.suffix.lstrip(".")
                    )
                    return self._normalize_pages(parsed["pages"]), parsed
            except (FileParseError, OSError, ValueError):
                pass
        if file_id != "main":
            try:
                attachment = task.attachment_texts[int(file_id.split(":", 1)[1]) - 1]
            except (IndexError, TypeError, ValueError):
                raise FileParseError("合同附件不存在。")
            pages = self.file_service._paginate_blocks(
                self.file_service._text_blocks(str(attachment.get("content", "")))
            )
            return self._normalize_pages(pages), attachment
        pages = self.file_service._paginate_blocks(
            self.file_service._text_blocks(task.contract_text)
        )
        return self._normalize_pages(pages), {}

    def _normalize_pages(self, pages: list[dict]) -> list[dict]:
        return self.file_service._document_result(
            pages, "not_required", False, None, False
        )["pages"]

    def _stored_location_is_valid(self, task: ReviewTask, risk: ReviewRisk) -> bool:
        if not risk.locatable or not risk.source_file_id:
            return False
        spans = list(risk.location_spans or [])
        if not spans and all(
            value is not None
            for value in (
                risk.page_number,
                risk.block_id,
                risk.start_offset,
                risk.end_offset,
                risk.anchor_text,
                risk.anchor_hash,
            )
        ):
            spans = [
                {
                    "source_file_id": risk.source_file_id,
                    "page_number": risk.page_number,
                    "block_id": risk.block_id,
                    "start_offset": risk.start_offset,
                    "end_offset": risk.end_offset,
                    "anchor_text": risk.anchor_text,
                    "anchor_hash": risk.anchor_hash,
                }
            ]
        if not spans:
            return False
        try:
            pages = self.pages(task, risk.source_file_id)
        except FileParseError:
            return False
        for span in spans:
            if span.get("source_file_id") != risk.source_file_id:
                return False
            page = next(
                (
                    item
                    for item in pages
                    if int(item.get("page_number", 0)) == span.get("page_number")
                ),
                None,
            )
            block = next(
                (
                    item
                    for item in (page or {}).get("blocks", [])
                    if item.get("block_id") == span.get("block_id")
                ),
                None,
            )
            if not block:
                return False
            text = str(block.get("text", ""))
            start = span.get("start_offset")
            end = span.get("end_offset")
            anchor = span.get("anchor_text")
            if not isinstance(start, int) or not isinstance(end, int):
                return False
            if not (0 <= start < end <= len(text)) or text[start:end] != anchor:
                return False
            if span.get("anchor_hash") != self._anchor_hash(
                risk.source_file_id,
                int(span["page_number"]),
                str(span["block_id"]),
                start,
                end,
                str(anchor),
            ):
                return False
        return True

    def _anchor_hash(
        self,
        source_file_id: str,
        page_number: int,
        block_id: str,
        start_offset: int,
        end_offset: int,
        anchor_text: str,
    ) -> str:
        return hashlib.sha256(
            (
                f"{source_file_id}|{page_number}|{block_id}|{start_offset}|{end_offset}|"
                f"{self._normalize(anchor_text)}"
            ).encode("utf-8")
        ).hexdigest()

    def _spans_from_references(
        self,
        source_file_id: str,
        references: list[tuple[int, str, int, str]],
    ) -> list[dict]:
        if not references:
            return []
        groups: list[list[tuple[int, str, int, str]]] = []
        for reference in references:
            previous = groups[-1][-1] if groups else None
            if previous and reference[:2] == previous[:2]:
                groups[-1].append(reference)
            else:
                groups.append([reference])
        spans: list[dict] = []
        for group in groups:
            page_number, block_id, start_offset, block_text = group[0]
            end_offset = group[-1][2] + 1
            anchor_text = block_text[start_offset:end_offset]
            spans.append(
                {
                    "source_file_id": source_file_id,
                    "page_number": page_number,
                    "block_id": block_id,
                    "start_offset": start_offset,
                    "end_offset": end_offset,
                    "anchor_text": anchor_text,
                    "anchor_hash": self._anchor_hash(
                        source_file_id,
                        page_number,
                        block_id,
                        start_offset,
                        end_offset,
                        anchor_text,
                    ),
                }
            )
        return spans

    @staticmethod
    def _normalized_with_offsets(value: str) -> list[tuple[str, int]]:
        punctuation = str.maketrans(
            {
                "，": ",",
                "。": ".",
                "；": ";",
                "：": ":",
                "（": "(",
                "）": ")",
                "【": "[",
                "】": "]",
                "“": '"',
                "”": '"',
                "‘": "'",
                "’": "'",
                "、": ",",
                "—": "-",
                "–": "-",
            }
        )
        normalized: list[tuple[str, int]] = []
        for offset, character in enumerate(value):
            mapped = unicodedata.normalize("NFKC", character).translate(punctuation)
            normalized.extend(
                (item, offset) for item in mapped if not item.isspace()
            )
        return normalized

    @classmethod
    def _normalized_value(cls, value: str) -> str:
        return "".join(item for item, _offset in cls._normalized_with_offsets(value))

    @classmethod
    def _matching_spans(cls, text: str, quote: str) -> list[tuple[int, int]]:
        normalized = cls._normalized_with_offsets(text)
        normalized_text = [item for item, _offset in normalized]
        offsets = [offset for _item, offset in normalized]
        normalized_quote = cls._normalized_value(quote)
        if not normalized_quote:
            return []
        haystack = "".join(normalized_text)
        spans: list[tuple[int, int]] = []
        cursor = 0
        while True:
            found = haystack.find(normalized_quote, cursor)
            if found < 0:
                break
            end_index = found + len(normalized_quote) - 1
            spans.append((offsets[found], offsets[end_index] + 1))
            cursor = found + max(1, len(normalized_quote))
        return spans

    @staticmethod
    def _source_file_id(task: ReviewTask, source: str | None) -> str | None:
        if source == task.contract_name:
            return "main"
        matches = [
            index
            for index, file_name in enumerate(task.attachment_names, start=1)
            if file_name == source
        ]
        return f"attachment:{matches[0]}" if len(matches) == 1 else None

    @staticmethod
    def _manifest_material(manifest: dict, file_id: str) -> dict:
        if file_id == "main":
            return manifest["main_file"]
        try:
            index = int(file_id.split(":", 1)[1]) - 1
            return manifest["attachments"][index]
        except (IndexError, KeyError, TypeError, ValueError) as exc:
            raise FileParseError("合同附件不存在。") from exc

    @staticmethod
    def _quad_points(rect: tuple[float, float, float, float]) -> ArrayObject:
        x1, y1, x2, y2 = rect
        return ArrayObject(
            [FloatObject(value) for value in (x1, y2, x2, y2, x1, y1, x2, y1)]
        )

    @staticmethod
    def _normalize(value: str) -> str:
        return DocumentService._normalized_value(value)

    @classmethod
    def _snapshot_matches_contract_text(
        cls, pages: list[dict], contract_text: str
    ) -> bool:
        if not pages or not contract_text:
            return False
        page_text = "".join(str(page.get("text", "")) for page in pages)
        return cls._normalize(page_text) == cls._normalize(contract_text)
