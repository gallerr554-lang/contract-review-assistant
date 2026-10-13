from pathlib import Path
import json

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from app.core.config import settings
from app.models.review import ReviewRisk, ReviewTask
from app.service.document_service import DocumentService
from app.service.file_service import FileParseError, FileService


def test_parse_utf8_txt(tmp_path: Path) -> None:
    contract = tmp_path / "contract.txt"
    contract.write_text("第一条 合同金额为人民币十万元。", encoding="utf-8")

    assert FileService(tmp_path).parse_file(contract) == "第一条 合同金额为人民币十万元。"


def test_manifest_paths_recover_after_storage_is_mounted_elsewhere(tmp_path: Path) -> None:
    service = FileService(tmp_path / "storage")
    upload_id = "5" * 36
    upload_dir = service.contract_root / upload_id
    upload_dir.mkdir(parents=True)
    source = upload_dir / "main.docx"
    preview = upload_dir / "main-preview.pdf"
    source.write_bytes(b"docx")
    preview.write_bytes(b"pdf")
    (upload_dir / "manifest.json").write_text(
        json.dumps(
            {
                "upload_id": upload_id,
                "user_id": 7,
                "main_file": {
                    "stored_path": "/old/host/storage/contracts/id/main.docx",
                    "preview_path": "/old/host/storage/contracts/id/main-preview.pdf",
                },
                "attachments": [],
            }
        ),
        encoding="utf-8",
    )

    manifest = service.load_manifest(upload_id, 7)

    assert manifest["main_file"]["stored_path"] == str(source)
    assert manifest["main_file"]["preview_path"] == str(preview)


@pytest.mark.parametrize("encoding", ["utf-16", "gb18030"])
def test_parse_common_chinese_txt_encodings(tmp_path: Path, encoding: str) -> None:
    contract = tmp_path / f"contract-{encoding}.txt"
    expected = "中文合同\n第一条 付款应在验收后十日内完成。"
    contract.write_bytes(expected.encode(encoding))

    assert FileService(tmp_path).parse_file(contract) == expected


def test_docx_keeps_headings_tables_pages_and_stable_block_ids(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    contract = tmp_path / "中文合同.docx"
    document = Document()
    document.add_heading("技术服务合同", level=1)
    document.add_paragraph("第一条 服务内容")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "系统开发"
    table.cell(1, 1).text = "人民币二十万元"
    document.add_page_break()
    document.add_heading("第二条 付款", level=1)
    document.add_paragraph("验收后十日内付款。")
    document.save(contract)
    monkeypatch.setattr(FileService, "_soffice_path", staticmethod(lambda: None))

    parsed = FileService(tmp_path / "storage").parse_document(contract)

    assert parsed["page_count"] == 2
    assert parsed["pages"][0]["blocks"][0]["style"].startswith("Heading")
    assert any(block["type"] == "table" for block in parsed["pages"][0]["blocks"])
    assert parsed["pages"][1]["text"].startswith("第二条 付款")
    assert all(
        block["block_id"].startswith(f"p{page['page_number']}:b")
        for page in parsed["pages"]
        for block in page["blocks"]
    )


def test_reject_unknown_extension(tmp_path: Path) -> None:
    contract = tmp_path / "contract.exe"
    contract.write_bytes(b"not a contract")

    with pytest.raises(FileParseError):
        FileService(tmp_path).parse_file(contract)


def test_scanned_pdf_uses_ocr_and_keeps_page_source(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    with source.open("wb") as output:
        writer.write(output)
    service = FileService(tmp_path / "storage")
    monkeypatch.setattr(service, "_ocr_pdf_page", lambda _path, _page: "扫描合同第一页")

    parsed = service.parse_document(source)

    assert parsed["ocr_status"] == "success"
    assert parsed["ocr_used"] is True
    assert parsed["text_source"] == "ocr"
    assert parsed["pages"][0]["ocr_used"] is True
    assert parsed["pages"][0]["text_source"] == "ocr"
    assert parsed["pages"][0]["text"] == "扫描合同第一页"


def test_repeated_cjk_glyph_pdf_text_is_not_treated_as_valid() -> None:
    damaged_text = "智" * 80 + " HYZN-CG-2026-0731-01"

    assert FileService._text_is_usable(damaged_text) is False


def test_long_text_is_paginated_without_rendering_single_document(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "document_page_characters", 20)
    contract = tmp_path / "long.txt"
    contract.write_text("第一段合同内容。\n第二段合同内容较长。\n第三段合同内容。", encoding="utf-8")

    parsed = FileService(tmp_path / "storage").parse_document(contract)

    assert parsed["page_count"] >= 2
    assert "第一段" in parsed["pages"][0]["text"]


def test_office_preview_uses_physical_pdf_page_count(
    tmp_path: Path,
) -> None:
    preview = tmp_path / "preview.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.add_blank_page(width=595, height=842)
    with preview.open("wb") as output:
        writer.write(output)
    blocks = [
        {"type": "paragraph", "text": "第一页合同内容"},
        {"type": "page_break", "text": ""},
        {"type": "paragraph", "text": "第二页合同内容"},
    ]

    pages = FileService._office_pages(blocks, preview)

    assert len(pages) == 2
    assert pages[0]["text"] == "第一页合同内容"
    assert pages[1]["text"] == "第二页合同内容"


def test_office_unusable_pdf_text_keeps_original_structured_blocks(
    tmp_path: Path,
) -> None:
    preview = tmp_path / "preview.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.add_blank_page(width=595, height=842)
    with preview.open("wb") as output:
        writer.write(output)
    blocks = [
        {"type": "paragraph", "text": f"第{index}条 中文合同内容"}
        for index in range(1, 9)
    ]

    pages = FileService._office_pages(blocks, preview)

    assert len(pages) == 2
    assert sum(len(page["blocks"]) for page in pages) == len(blocks)
    assert "中文合同内容" in "".join(page["text"] for page in pages)
    assert all(block["type"] == "paragraph" for page in pages for block in page["blocks"])
    assert FileService._office_preview_is_usable(preview) is False


def test_office_conversion_replaces_unusable_cached_preview(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "中文合同.docx"
    source.write_bytes(b"source")
    preview = tmp_path / "中文合同-preview.pdf"
    preview.write_bytes(b"damaged-preview")
    conversions: list[Path] = []

    def fake_conversion(path: Path, output_dir: Path, _soffice: str) -> Path:
        conversions.append(path)
        converted = output_dir / "中文合同.pdf"
        converted.write_bytes(b"valid-preview")
        return converted

    monkeypatch.setattr(FileService, "_soffice_path", staticmethod(lambda: "/fake/soffice"))
    monkeypatch.setattr(
        FileService,
        "_office_preview_is_usable",
        staticmethod(lambda path: path.read_bytes() == b"valid-preview"),
    )
    monkeypatch.setattr(FileService, "_run_office_conversion", staticmethod(fake_conversion))

    converted = FileService._convert_office_to_pdf(source)

    assert converted == preview
    assert preview.read_bytes() == b"valid-preview"
    assert conversions == [source]


def test_macos_office_conversion_uses_system_fontconfig(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr("app.service.file_service.sys.platform", "darwin")

    environment = FileService._office_conversion_environment(
        tmp_path, "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    )

    config_path = Path(environment["FONTCONFIG_FILE"])
    config = config_path.read_text(encoding="utf-8")
    assert environment["XDG_CACHE_HOME"] == str(tmp_path / "font-cache")
    assert "/System/Library/Fonts" in config
    assert "/System/Library/Fonts/Supplemental" in config
    assert '<cachedir prefix="xdg">fontconfig</cachedir>' in config


def test_historical_office_snapshot_is_rebuilt_from_physical_pages(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    storage = tmp_path / "storage"
    file_service = FileService(storage)
    upload_id = "2" * 36
    upload_dir = file_service.contract_root / upload_id
    upload_dir.mkdir(parents=True)
    source = upload_dir / "main.docx"
    source.write_bytes(b"historical-office-placeholder")
    (upload_dir / "manifest.json").write_text(
        json.dumps(
            {
                "upload_id": upload_id,
                "user_id": 1,
                "main_file": {
                    "stored_path": str(source),
                    "file_name": "历史合同.docx",
                    "file_extension": "docx",
                    "pages": [{"page_number": 1, "text": "旧字符分页", "blocks": []}],
                },
                "attachments": [],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        file_service,
        "parse_document",
        lambda _path, _extension: {
            "pages": [
                {"page_number": 1, "text": "第一页", "blocks": []},
                {"page_number": 2, "text": "第二页", "blocks": []},
            ],
            "page_count": 2,
            "pagination_source": "physical_pdf",
            "preview_available": True,
        },
    )
    task = ReviewTask(
        id=94,
        task_no="CR-HISTORICAL-OFFICE",
        user_id=1,
        upload_id=upload_id,
        contract_name="历史合同.docx",
        contract_text="第一页\n第二页",
        attachment_names=[],
        attachment_texts=[],
        scene_name="测试",
        scene_snapshot={"rules": []},
        model_snapshot={},
        document_pages=[{"page_number": 1, "text": "旧字符分页", "blocks": []}],
        document_meta={"page_count": 1, "office_text_source": "structured"},
        scale="标准",
        position="甲方",
        status="SUCCESS",
        progress=100,
        current_step=5,
        current_message="完成",
    )

    changed = DocumentService(file_service).ensure_document_snapshot(task)

    assert changed is True
    assert [page["text"] for page in task.document_pages] == ["第一页", "第二页"]
    assert task.document_meta["page_count"] == 2
    assert task.document_meta["pagination_source"] == "physical_pdf"
    assert task.document_meta["office_text_source"] == "structured"


def test_historical_office_snapshot_rebuilds_unusable_preview(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    storage = tmp_path / "storage"
    file_service = FileService(storage)
    upload_id = "4" * 36
    upload_dir = file_service.contract_root / upload_id
    upload_dir.mkdir(parents=True)
    source = upload_dir / "main.docx"
    source.write_bytes(b"historical-office-placeholder")
    (upload_dir / "manifest.json").write_text(
        json.dumps(
            {
                "upload_id": upload_id,
                "user_id": 1,
                "main_file": {
                    "stored_path": str(source),
                    "file_name": "历史合同.docx",
                    "file_extension": "docx",
                },
                "attachments": [],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        file_service,
        "parse_document",
        lambda _path, _extension: {
            "pages": [
                {"page_number": 1, "text": "第一页", "blocks": []},
                {"page_number": 2, "text": "第二页", "blocks": []},
            ],
            "page_count": 2,
            "pagination_source": "physical_pdf",
            "office_text_source": "structured",
            "preview_quality": "usable",
            "preview_available": True,
        },
    )
    task = ReviewTask(
        id=95,
        task_no="CR-HISTORICAL-BAD-PREVIEW",
        user_id=1,
        upload_id=upload_id,
        contract_name="历史合同.docx",
        contract_text="旧分页文本",
        attachment_names=[],
        attachment_texts=[],
        scene_name="测试",
        scene_snapshot={"rules": []},
        model_snapshot={},
        document_pages=[{"page_number": 1, "text": "旧分页文本", "blocks": []}],
        document_meta={
            "page_count": 1,
            "office_text_source": "structured",
            "preview_quality": "unusable",
        },
        scale="标准",
        position="甲方",
        status="SUCCESS",
        progress=100,
        current_step=5,
        current_message="完成",
    )

    changed = DocumentService(file_service).ensure_document_snapshot(task)

    assert changed is True
    assert [page["text"] for page in task.document_pages] == ["第一页", "第二页"]
    assert task.document_meta["page_count"] == 2
    assert task.document_meta["preview_quality"] == "usable"


def test_historical_pdf_damaged_text_is_reparsed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    storage = tmp_path / "storage"
    file_service = FileService(storage)
    upload_id = "3" * 36
    upload_dir = file_service.contract_root / upload_id
    upload_dir.mkdir(parents=True)
    source = upload_dir / "main.pdf"
    source.write_bytes(b"historical-pdf-placeholder")
    (upload_dir / "manifest.json").write_text(
        json.dumps(
            {
                "upload_id": upload_id,
                "user_id": 1,
                "main_file": {
                    "stored_path": str(source),
                    "file_name": "历史合同.pdf",
                    "file_extension": "pdf",
                    "pages": [{"page_number": 1, "text": "智" * 80, "blocks": []}],
                },
                "attachments": [],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        file_service,
        "parse_document",
        lambda _path, _extension: {
            "pages": [{"page_number": 1, "text": "重新识别的中文合同", "blocks": []}],
            "page_count": 1,
            "pagination_source": "physical_pdf",
            "ocr_status": "success",
            "ocr_used": True,
            "text_source": "ocr",
            "preview_available": True,
        },
    )
    task = ReviewTask(
        id=95,
        task_no="CR-HISTORICAL-PDF",
        user_id=1,
        upload_id=upload_id,
        contract_name="历史合同.pdf",
        contract_text="重新识别的中文合同",
        attachment_names=[],
        attachment_texts=[],
        scene_name="测试",
        scene_snapshot={"rules": []},
        model_snapshot={},
        document_pages=[{"page_number": 1, "text": "智" * 80, "blocks": []}],
        document_meta={"page_count": 1, "ocr_status": "not_required"},
        scale="标准",
        position="甲方",
        status="SUCCESS",
        progress=100,
        current_step=5,
        current_message="完成",
    )

    changed = DocumentService(file_service).ensure_document_snapshot(task)

    assert changed is True
    assert task.document_pages[0]["text"] == "重新识别的中文合同"
    assert task.document_meta["ocr_status"] == "success"
    assert task.document_meta["text_source"] == "ocr"


def test_marked_pdf_preserves_original_page_and_adds_annotation(tmp_path: Path) -> None:
    storage = tmp_path / "storage"
    file_service = FileService(storage)
    upload_id = "c" * 36
    upload_dir = file_service.contract_root / upload_id
    upload_dir.mkdir(parents=True)
    source = upload_dir / "main.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    with source.open("wb") as output:
        writer.write(output)
    (upload_dir / "manifest.json").write_text(
        json.dumps(
            {
                "upload_id": upload_id,
                "user_id": 9,
                "main_file": {"stored_path": str(source), "file_name": "合同.pdf"},
                "attachments": [],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    task = ReviewTask(
        id=88,
        task_no="CR-PDF",
        user_id=9,
        upload_id=upload_id,
        contract_name="合同.pdf",
        contract_text="付款节点另行约定",
        attachment_names=[],
        attachment_texts=[],
        scene_name="测试",
        scene_snapshot={"rules": []},
        model_snapshot={},
        document_pages=[{"page_number": 1, "text": "付款节点另行约定", "blocks": []}],
        document_meta={"page_count": 1},
        scale="标准",
        position="甲方",
        status="SUCCESS",
        progress=100,
        current_step=5,
        current_message="完成",
    )
    risk = ReviewRisk(
        id=5,
        task_id=88,
        rule_id="R003",
        rule_name="付款审查",
        risk_name="付款节点不明确",
        risk_level="high",
        original_text="付款节点另行约定",
        source="合同.pdf",
        locatable=True,
        description="付款条件不清晰。",
        suggestion="明确付款节点。",
        page_number=1,
    )

    output = DocumentService(file_service).generate_marked_pdf(task, 9, [risk])
    marked = PdfReader(output)

    assert len(marked.pages) == 1
    assert len(marked.pages[0].get("/Annots", [])) >= 1


def test_marked_pdf_keeps_unlocatable_risk_as_first_page_note(tmp_path: Path) -> None:
    storage = tmp_path / "storage"
    file_service = FileService(storage)
    upload_id = "d" * 36
    upload_dir = file_service.contract_root / upload_id
    upload_dir.mkdir(parents=True)
    source = upload_dir / "main.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    with source.open("wb") as output:
        writer.write(output)
    (upload_dir / "manifest.json").write_text(
        json.dumps(
            {
                "upload_id": upload_id,
                "user_id": 10,
                "main_file": {"stored_path": str(source), "file_name": "合同.pdf"},
                "attachments": [],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    task = ReviewTask(
        id=89,
        task_no="CR-PDF-UNLOCATED",
        user_id=10,
        upload_id=upload_id,
        contract_name="合同.pdf",
        contract_text="合同正文",
        attachment_names=[],
        attachment_texts=[],
        scene_name="测试",
        scene_snapshot={"rules": []},
        model_snapshot={},
        document_pages=[{"page_number": 1, "text": "合同正文", "blocks": []}],
        document_meta={"page_count": 1},
        scale="标准",
        position="甲方",
        status="SUCCESS",
        progress=100,
        current_step=5,
        current_message="完成",
    )
    risk = ReviewRisk(
        id=6,
        task_id=89,
        rule_id="R004",
        rule_name="附件审查",
        risk_name="附件条款需核验",
        risk_level="medium",
        original_text="附件中的原文",
        source="附件.docx",
        locatable=False,
        description="风险说明。",
        suggestion="人工核验附件。",
        page_number=None,
    )

    output = DocumentService(file_service).generate_marked_pdf(task, 10, [risk])
    marked = PdfReader(output)

    assert len(marked.pages[0].get("/Annots", [])) == 1


def test_marked_docx_preserves_content_adds_comments_and_manual_section(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    storage = tmp_path / "storage"
    file_service = FileService(storage)
    upload_id = "9" * 36
    upload_dir = file_service.contract_root / upload_id
    upload_dir.mkdir(parents=True)
    source = upload_dir / "main.docx"
    source_document = Document()
    source_document.add_heading("设备采购合同", level=1)
    source_document.add_paragraph("验收合格后十日内支付全部价款。")
    table = source_document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "付款节点"
    table.cell(0, 1).text = "付款比例"
    table.cell(1, 0).text = "验收后"
    table.cell(1, 1).text = "100%"
    source_document.save(source)
    monkeypatch.setattr(FileService, "_soffice_path", staticmethod(lambda: None))
    parsed = file_service.parse_document(source)
    (upload_dir / "manifest.json").write_text(
        json.dumps(
            {
                "upload_id": upload_id,
                "user_id": 12,
                "main_file": {
                    "stored_path": str(source),
                    "file_name": "设备采购合同.docx",
                    "file_extension": "docx",
                    **parsed,
                },
                "attachments": [],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    task = ReviewTask(
        id=98,
        task_no="CR-MARKED-DOCX",
        user_id=12,
        upload_id=upload_id,
        contract_name="设备采购合同.docx",
        contract_text=parsed["content"],
        attachment_names=[],
        attachment_texts=[],
        scene_name="测试",
        scene_snapshot={"rules": []},
        model_snapshot={},
        document_pages=parsed["pages"],
        document_meta={"file_name": "设备采购合同.docx", "file_extension": "docx"},
        document_files=[
            DocumentService.snapshot_file(
                "main",
                {"file_name": "设备采购合同.docx", "file_extension": "docx", **parsed},
            )
        ],
        scale="标准",
        position="甲方",
        status="SUCCESS",
        progress=100,
        current_step=5,
        current_message="完成",
    )
    service = DocumentService(file_service)
    location = service.locate_text(
        task, task.contract_name, "验收合格后十日内支付全部价款"
    )
    assert location is not None
    primary = location["location_spans"][0]
    located = ReviewRisk(
        id=21,
        task_id=98,
        rule_id="R003",
        rule_name="付款审查",
        risk_name="付款期限",
        risk_level="high",
        original_text="验收合格后十日内支付全部价款",
        source=task.contract_name,
        locatable=True,
        description="付款条件需要核验。",
        suggestion="明确付款前置条件。",
        page_number=primary["page_number"],
        source_file_id="main",
        block_id=primary["block_id"],
        start_offset=primary["start_offset"],
        end_offset=primary["end_offset"],
        anchor_text=primary["anchor_text"],
        anchor_hash=primary["anchor_hash"],
        location_spans=location["location_spans"],
    )
    manual = ReviewRisk(
        id=22,
        task_id=98,
        rule_id="R005",
        rule_name="附件审查",
        risk_name="附件缺失",
        risk_level="medium",
        original_text=None,
        source="技术附件.docx",
        locatable=False,
        description="附件内容需核验。",
        suggestion="补充附件后复核。",
    )

    output = service.generate_marked_docx(task, 12, [located, manual])
    marked = Document(output)

    assert output.suffix == ".docx"
    assert marked.paragraphs[0].text == "设备采购合同"
    assert any(
        run.font.highlight_color
        for paragraph in marked.paragraphs
        for run in paragraph.runs
    )
    assert len(marked.comments) == 1
    assert "付款期限" in next(iter(marked.comments)).text
    assert any(paragraph.text == "人工核验事项" for paragraph in marked.paragraphs)
    assert any(
        cell.text == "附件缺失"
        for marked_table in marked.tables
        for row in marked_table.rows
        for cell in row.cells
    )


def test_risk_location_uses_unique_stable_block_anchor(tmp_path: Path) -> None:
    task = ReviewTask(
        id=91,
        task_no="CR-LOCATION",
        user_id=1,
        upload_id="e" * 36,
        contract_name="中文合同.txt",
        contract_text="第一条 付款应在验收后十日内完成。",
        attachment_names=[],
        attachment_texts=[],
        scene_name="测试",
        scene_snapshot={"rules": []},
        model_snapshot={},
        document_pages=[
            {
                "page_number": 1,
                "text": "第一条 付款应在验收后十日内完成。",
                "blocks": [
                    {
                        "block_id": "p1:b1",
                        "type": "paragraph",
                        "text": "第一条 付款应在验收后十日内完成。",
                    }
                ],
            }
        ],
        document_meta={"page_count": 1},
        scale="标准",
        position="甲方",
        status="SUCCESS",
        progress=100,
        current_step=5,
        current_message="完成",
    )
    service = DocumentService(FileService(tmp_path / "storage"))

    location = service.locate_text(task, "中文合同.txt", "付款应在验收后十日内完成")

    assert location is not None
    assert location["source_file_id"] == "main"
    assert location["page_number"] == 1
    assert location["block_id"] == "p1:b1"
    assert location["start_offset"] < location["end_offset"]
    assert location["anchor_text"] == "付款应在验收后十日内完成"
    assert len(location["anchor_hash"]) == 64


def test_unique_quote_ignores_incorrect_model_page_hint(tmp_path: Path) -> None:
    pages = [
        {
            "page_number": 1,
            "text": "合同首页。",
            "blocks": [{"block_id": "p1:b1", "type": "paragraph", "text": "合同首页。"}],
        },
        {
            "page_number": 2,
            "text": "验收合格后十日内支付全部价款。",
            "blocks": [{"block_id": "p2:b1", "type": "paragraph", "text": "验收合格后十日内支付全部价款。"}],
        },
    ]
    task = ReviewTask(
        id=97,
        task_no="CR-WRONG-PAGE-HINT",
        user_id=1,
        upload_id="7" * 36,
        contract_name="分页合同.docx",
        contract_text="\n".join(page["text"] for page in pages),
        attachment_names=[],
        attachment_texts=[],
        scene_name="测试",
        scene_snapshot={"rules": []},
        model_snapshot={},
        document_pages=pages,
        document_meta={"page_count": 2},
        scale="标准",
        position="甲方",
        status="SUCCESS",
        progress=100,
        current_step=5,
        current_message="完成",
    )

    location = DocumentService(FileService(tmp_path / "storage")).locate_text(
        task,
        task.contract_name,
        "验收合格后十日内支付全部价款",
        page_number=5,
    )

    assert location is not None
    assert location["page_number"] == 2


def test_duplicate_original_text_requires_manual_verification(tmp_path: Path) -> None:
    repeated = "双方另行协商。"
    task = ReviewTask(
        id=92,
        task_no="CR-DUPLICATE",
        user_id=1,
        upload_id="f" * 36,
        contract_name="重复条款.txt",
        contract_text=f"付款：{repeated}\n交付：{repeated}",
        attachment_names=[],
        attachment_texts=[],
        scene_name="测试",
        scene_snapshot={"rules": []},
        model_snapshot={},
        document_pages=[
            {
                "page_number": 1,
                "text": f"付款：{repeated}\n交付：{repeated}",
                "blocks": [
                    {"block_id": "p1:b1", "type": "paragraph", "text": f"付款：{repeated}"},
                    {"block_id": "p1:b2", "type": "paragraph", "text": f"交付：{repeated}"},
                ],
            }
        ],
        document_meta={"page_count": 1},
        scale="标准",
        position="甲方",
        status="SUCCESS",
        progress=100,
        current_step=5,
        current_message="完成",
    )

    location = DocumentService(FileService(tmp_path / "storage")).locate_text(
        task, "重复条款.txt", repeated
    )

    assert location is None


def test_duplicate_original_text_uses_exact_context_and_page_hint(tmp_path: Path) -> None:
    repeated = "双方另行协商。"
    pages = [
        {
            "page_number": 1,
            "text": f"付款安排：{repeated}",
            "blocks": [{"block_id": "p1:b1", "type": "paragraph", "text": f"付款安排：{repeated}"}],
        },
        {
            "page_number": 2,
            "text": f"交付安排：{repeated}具体日期另行通知。",
            "blocks": [{"block_id": "p2:b1", "type": "paragraph", "text": f"交付安排：{repeated}具体日期另行通知。"}],
        },
    ]
    task = ReviewTask(
        id=94,
        task_no="CR-DUPLICATE-CONTEXT",
        user_id=1,
        upload_id="2" * 36,
        contract_name="重复条款.txt",
        contract_text="\n".join(page["text"] for page in pages),
        attachment_names=[],
        attachment_texts=[],
        scene_name="测试",
        scene_snapshot={"rules": []},
        model_snapshot={},
        document_pages=pages,
        document_meta={"page_count": 2},
        scale="标准",
        position="甲方",
        status="SUCCESS",
        progress=100,
        current_step=5,
        current_message="完成",
    )
    service = DocumentService(FileService(tmp_path / "storage"))

    location = service.locate_text(
        task,
        task.contract_name,
        repeated,
        page_number=2,
        context_before="交付安排：",
        context_after="具体日期另行通知。",
    )

    assert location is not None
    assert location["page_number"] == 2
    assert location["block_id"] == "p2:b1"
    assert service.locate_text(
        task, task.contract_name, repeated, context_before="不存在的上下文"
    ) is None


def test_risk_quote_can_span_adjacent_blocks_and_table_cells(tmp_path: Path) -> None:
    pages = [
        {
            "page_number": 1,
            "text": "付款条件如下：\n验收合格后十日内支付。\n付款节点 | 比例 | 验收后支付 | 100%",
            "blocks": [
                {"block_id": "p1:b1", "type": "paragraph", "text": "付款条件如下："},
                {"block_id": "p1:b2", "type": "paragraph", "text": "验收合格后十日内支付。"},
                {
                    "block_id": "p1:b3",
                    "type": "table",
                    "text": "付款节点 | 比例\n验收后支付 | 100%",
                    "rows": [["付款节点", "比例"], ["验收后支付", "100%"]],
                },
            ],
        }
    ]
    task = ReviewTask(
        id=95,
        task_no="CR-MULTI-BLOCK",
        user_id=1,
        upload_id="3" * 36,
        contract_name="跨段合同.docx",
        contract_text=pages[0]["text"],
        attachment_names=[],
        attachment_texts=[],
        scene_name="测试",
        scene_snapshot={"rules": []},
        model_snapshot={},
        document_pages=pages,
        document_meta={"page_count": 1},
        scale="标准",
        position="甲方",
        status="SUCCESS",
        progress=100,
        current_step=5,
        current_message="完成",
    )
    service = DocumentService(FileService(tmp_path / "storage"))

    paragraph_location = service.locate_text(
        task, task.contract_name, "付款条件如下：验收合格后十日内支付。"
    )
    table_location = service.locate_text(
        task, task.contract_name, "付款节点 | 比例\n验收后支付 | 100%"
    )

    assert paragraph_location is not None
    assert [item["block_id"] for item in paragraph_location["location_spans"]] == [
        "p1:b1",
        "p1:b2",
    ]
    assert table_location is not None
    assert len(table_location["location_spans"]) == 1
    assert table_location["location_spans"][0]["block_id"] == "p1:b3"


def test_attachment_risk_uses_attachment_snapshot(tmp_path: Path) -> None:
    attachment_pages = [
        {
            "page_number": 1,
            "text": "技术附件第一条。",
            "blocks": [{"block_id": "p1:b1", "type": "paragraph", "text": "技术附件第一条。"}],
        },
        {
            "page_number": 2,
            "text": "设备验收不合格时，乙方应在五日内完成整改。",
            "blocks": [{"block_id": "p2:b1", "type": "paragraph", "text": "设备验收不合格时，乙方应在五日内完成整改。"}],
        },
    ]
    task = ReviewTask(
        id=96,
        task_no="CR-ATTACHMENT",
        user_id=1,
        upload_id="4" * 36,
        contract_name="主合同.txt",
        contract_text="主合同正文。",
        attachment_names=["技术附件.docx"],
        attachment_texts=[{"file_name": "技术附件.docx", "content": "\n".join(page["text"] for page in attachment_pages)}],
        scene_name="测试",
        scene_snapshot={"rules": []},
        model_snapshot={},
        document_pages=[{"page_number": 1, "text": "主合同正文。", "blocks": [{"block_id": "p1:b1", "type": "paragraph", "text": "主合同正文。"}]}],
        document_meta={"file_name": "主合同.txt", "page_count": 1},
        document_files=[
            {"file_id": "main", "file_name": "主合同.txt", "file_extension": "txt", "pages": [{"page_number": 1, "text": "主合同正文。", "blocks": [{"block_id": "p1:b1", "type": "paragraph", "text": "主合同正文。"}]}]},
            {"file_id": "attachment:1", "file_name": "技术附件.docx", "file_extension": "docx", "pages": attachment_pages},
        ],
        scale="标准",
        position="甲方",
        status="SUCCESS",
        progress=100,
        current_step=5,
        current_message="完成",
    )

    location = DocumentService(FileService(tmp_path / "storage")).locate_text(
        task,
        "技术附件.docx",
        "乙方应在五日内完成整改",
        source_file_id="attachment:1",
        page_number=2,
    )

    assert location is not None
    assert location["source_file_id"] == "attachment:1"
    assert location["page_number"] == 2
    assert all(
        item["source_file_id"] == "attachment:1"
        for item in location["location_spans"]
    )


def test_tampered_anchor_hash_is_rebuilt(tmp_path: Path) -> None:
    task = ReviewTask(
        id=93,
        task_no="CR-HASH",
        user_id=1,
        upload_id="1" * 36,
        contract_name="中文合同.txt",
        contract_text="第一条 验收后十日内付款。",
        attachment_names=[],
        attachment_texts=[],
        scene_name="测试",
        scene_snapshot={"rules": []},
        model_snapshot={},
        document_pages=[
            {
                "page_number": 1,
                "text": "第一条 验收后十日内付款。",
                "blocks": [
                    {
                        "block_id": "p1:b1",
                        "type": "paragraph",
                        "text": "第一条 验收后十日内付款。",
                    }
                ],
            }
        ],
        document_meta={"page_count": 1},
        scale="标准",
        position="甲方",
        status="SUCCESS",
        progress=100,
        current_step=5,
        current_message="完成",
    )
    risk = ReviewRisk(
        id=7,
        task_id=93,
        rule_id="R001",
        rule_name="付款审查",
        risk_name="付款期限",
        risk_level="medium",
        original_text="验收后十日内付款",
        source="中文合同.txt",
        locatable=True,
        description="付款期限需核验。",
        suggestion="确认付款条件。",
        page_number=1,
        source_file_id="main",
        block_id="p1:b1",
        start_offset=4,
        end_offset=13,
        anchor_text="验收后十日内付款",
        anchor_hash="0" * 64,
    )
    service = DocumentService(FileService(tmp_path / "storage"))

    changed = service.ensure_risk_locations(task, [risk])

    assert changed is True
    assert risk.locatable is True
    assert risk.anchor_hash != "0" * 64
    assert len(risk.anchor_hash or "") == 64
